# -*- coding: utf-8 -*-
"""Saídas dos fundos de previdência (HGD30, HYD60): Informativo e e-mail.

INFORMATIVO (PPTX + PDF)
  O modelo é o PPTX que vai para os clientes (templates/informativos/), e não
  um desenho novo: aqui só se trocam os números, o mês e o texto do gestor,
  com a formatação de cada caixa preservada. Gráficos nativos (rating,
  composição) recebem os dados novos; a imagem de setores é redesenhada no
  mesmo estilo. O HYD60 ganha a página de rentabilidade (tabela + histórico).
  O PDF sai do próprio PowerPoint (Windows com Office); sem ele, só o PPTX.

      saida/AAAA-MM/informativos/Informativo - HGD30 - Setembro 2026.pptx / .pdf

E-MAIL
  O Chromium abre o gerador (central/tivio-email-previdencia.html) e fotografa
  cada bloco: são os PNG que vão no Mailchimp. Ao lado, o e-mail em HTML (com
  os PNG ao lado), um .eml com as imagens embutidas (abre no Outlook pronto
  para enviar) e, com o Outlook instalado, o modelo .oft.

      saida/AAAA-MM/emails/previdencia/<fundo>/01-cabecalho.png …
      saida/AAAA-MM/emails/previdencia/<fundo>/email-<fundo>.html | .eml | .oft
"""
import copy
import io
import os
import re
import subprocess

from calculators import formatos as fmt
from exporters.pdf import achar_chromium

RAIZ = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
MESES_RE = '|'.join(fmt.MESES)

# Onde fica cada dado em cada modelo: (slide, nome do shape). Os nomes vêm do
# PPTX que o time usa; se o modelo mudar, é aqui que se ajusta.
MAPAS = {
    'hgd30': {
        'texto': (2, 'CaixaDeTexto 4'),
        'pl': (3, 'object 5'),
        'duration': (3, 'Rectangle: Rounded Corners 113'),
        'alocacao': (5, ['TextBox 54', 'TextBox 51', 'TextBox 48']),
        'rating': (5, 'Chart 6'),
        'composicao': (5, 'Chart 10'),
        'setorial': (5, 'Imagem 20'),
    },
    'hyd60': {
        'texto': (2, 'CaixaDeTexto 1'),
        'alocacao': (5, ['TextBox 54', 'TextBox 51', 'TextBox 48']),
        'rating': (5, 'Chart 6'),
        'setorial': (5, 'Imagem 6'),
        'pagina_rentabilidade': 5,        # entra logo depois desta página
    },
}

A = '{http://schemas.openxmlformats.org/drawingml/2006/main}'


def _shapes(colecao):
    for sh in colecao:
        if sh.shape_type == 6:            # grupo
            yield from _shapes(sh.shapes)
        else:
            yield sh


def _acha(slide, nome, filtro=None):
    achados = [s for s in _shapes(slide.shapes) if s.name == nome]
    if filtro:
        achados = [s for s in achados if filtro(s)]
    return achados


def _troca_runs(paragrafo, novos):
    """Põe `novos` (lista de textos) nos runs existentes, mantendo a fonte de
    cada um; o que sobra de runs fica vazio."""
    runs = paragrafo.runs
    for i, r in enumerate(runs):
        r.text = novos[i] if i < len(novos) else ''
    if len(novos) > len(runs) and runs:
        runs[-1].text = runs[-1].text + ''.join(novos[len(runs):])


# --------------------------------------------------------------- mês e data
def _mes_e_data(prs, d):
    for slide in prs.slides:
        for sh in _shapes(slide.shapes):
            if not sh.has_text_frame:
                continue
            for p in sh.text_frame.paragraphs:
                t = ''.join(r.text for r in p.runs)
                if re.fullmatch(rf'\s*({MESES_RE}) de \d{{4}}\s*', t):
                    # "Setembro de 2026" é mais longo que "Agosto de 2026": não quebra
                    sh.text_frame.word_wrap = False
                    if len(p.runs) >= 2 and p.runs[1].text.strip().startswith('de'):
                        _troca_runs(p, [d['mes'], f' de {d["ano"]}'])
                    else:
                        _troca_runs(p, [f'{d["mes"]} de {d["ano"]}'])
                for r in p.runs:
                    if re.search(r'Data base:\s*\d{2}/\d{2}/\d{4}', r.text):
                        r.text = re.sub(r'\d{2}/\d{2}/\d{4}', d['data_base'], r.text)


# ----------------------------------------------------------- texto do gestor
def _runs_markdown(texto):
    """'a **b** c' -> [('a ', False), ('b', True), (' c', False)]"""
    partes, negrito = [], False
    for pedaco in re.split(r'(\*\*)', texto):
        if pedaco == '**':
            negrito = not negrito
        elif pedaco:
            partes.append((pedaco, negrito))
    return partes


def _blocos_markdown(md):
    """[('h', texto) | ('li', texto) | ('p', texto)] na ordem do texto."""
    saida = []
    for bloco in re.split(r'\n\s*\n', md or ''):
        linhas = [l.strip() for l in bloco.splitlines() if l.strip()]
        if not linhas:
            continue
        if all(l.startswith('- ') for l in linhas):
            saida += [('li', l[2:]) for l in linhas]
        else:
            t = ' '.join(linhas)
            saida.append(('h', t.strip('*')) if re.fullmatch(r'\*\*[^*]+\*\*', t) else ('p', t))
    return saida


def _escreve_texto(shape, md):
    """Reescreve a caixa do texto do gestor com os parágrafos-modelo dela:
    subtítulo, item de lista, parágrafo e linha em branco."""
    tf = shape.text_frame
    ps = list(tf.paragraphs)

    def tem_bullet(p):
        return p._p.pPr is not None and p._p.pPr.find(A + 'buChar') is not None

    com_texto = [p for p in ps if p.runs and ''.join(r.text for r in p.runs).strip()]
    m_li = next((p for p in com_texto if tem_bullet(p)), None)
    m_h = next((p for p in com_texto if not tem_bullet(p) and all(r.font.bold for r in p.runs if r.text.strip())), None)
    m_p = next((p for p in com_texto if not tem_bullet(p) and p is not m_h), None) or m_h or com_texto[0]
    m_h = m_h or m_p
    m_li = m_li or m_p
    vazio = next((p for p in ps if not ''.join(r.text for r in p.runs).strip()), None)

    modelos = {k: copy.deepcopy(v._p) for k, v in (('h', m_h), ('li', m_li), ('p', m_p))}
    m_vazio = copy.deepcopy(vazio._p) if vazio is not None else None
    txBody = tf._txBody
    for p in ps:
        txBody.remove(p._p)

    def novo(tipo, texto):
        el = copy.deepcopy(modelos[tipo])
        rs = el.findall(A + 'r')
        base = copy.deepcopy(rs[0]) if rs else None
        for r in rs:
            el.remove(r)
        for extra in el.findall(A + 'br'):
            el.remove(extra)
        fim = el.find(A + 'endParaRPr')
        partes = [(texto, True)] if tipo == 'h' else _runs_markdown(texto)
        for t, b in partes:
            r = copy.deepcopy(base)
            rpr = r.find(A + 'rPr')
            if rpr is not None:
                rpr.set('b', '1' if b else '0')
                if rpr.find(A + 'latin') is None:
                    from lxml import etree
                    lat = etree.SubElement(rpr, A + 'latin', typeface='Versos')
                    # a ordem do schema: latin vem depois de preenchimento/efeitos
                    for tag in ('ea', 'cs', 'sym', 'hlinkClick', 'hlinkMouseOver', 'rtl', 'extLst'):
                        x = rpr.find(A + tag)
                        if x is not None:
                            x.addprevious(lat)
                            break
            r.find(A + 't').text = t
            if fim is not None:
                fim.addprevious(r)
            else:
                el.append(r)
        txBody.append(el)

    def espaco():
        if m_vazio is not None:
            txBody.append(copy.deepcopy(m_vazio))

    anterior = None
    for tipo, texto in _blocos_markdown(md):
        if anterior is not None and not (tipo == 'li' and anterior == 'li'):
            espaco()
        novo(tipo, texto)
        anterior = tipo
    _cabe(shape)


def _cabe(shape, largura_char=0.47, entrelinha=1.2):
    """Encolhe a letra da caixa até o texto caber na altura dela.

    O texto do gestor muda de tamanho todo mês e a caixa do modelo é fixa; o
    PowerPoint não refaz o "reduzir texto" ao abrir nem ao gerar o PDF. A
    medida é aproximada (largura média do caractere = 0,47 do corpo, calibrada
    no Informativo do HYD60) e só reduz, nunca aumenta.
    """
    import math
    tf = shape.text_frame
    W, H = shape.width / 12700, shape.height / 12700
    paras = []
    base = 11.0
    for p in tf.paragraphs:
        t = ''.join(r.text for r in p.runs)
        sz = next((r.font.size.pt for r in p.runs if r.font.size), None)
        if sz:
            base = sz
        marl = int(p._p.pPr.get('marL', '0')) / 12700 if p._p.pPr is not None else 0
        paras.append((t, marl))

    def altura(sz):
        linhas = 0
        for t, marl in paras:
            cpl = max(10, (W - marl) / (largura_char * sz))
            linhas += max(1, math.ceil(len(t) / cpl))
        return linhas * sz * entrelinha

    if altura(base) <= H * 0.98:
        return
    s = 1.0
    while s > 0.6 and altura(base * s) > H * 0.98:
        s -= 0.02
    novo = int(round(base * s * 100))
    for p in tf.paragraphs:
        for el in p._p.iter():
            if el.tag in (A + 'rPr', A + 'endParaRPr', A + 'defRPr'):
                el.set('sz', str(novo))


# ------------------------------------------------------------------ números
def _pl(shape, d):
    p = shape.text_frame.paragraphs[0]
    _troca_runs(p, ['R$ ', d['pl_curto'][3:]])


def _duration(shape, d):
    p = [p for p in shape.text_frame.paragraphs if p.runs][-1]
    _troca_runs(p, [d['duration_num'], ' ', 'anos' if (d['duration'] or 0) >= 1 else 'ano'])


def _alocacao(slide, nomes, d):
    alvo = [d['aloc_mapa'].get(n, fmt.MINUS) for n in ('Crédito High Grade', 'Crédito High Yield', 'Caixa')]
    for nome, valor in zip(nomes, alvo):
        for sh in _acha(slide, nome, lambda s: s.has_text_frame and '%' in s.text_frame.text
                        and 'CDI' not in s.text_frame.text):
            _troca_runs(sh.text_frame.paragraphs[0], [valor[:-1], '%'])


def _grafico(shape, categorias, valores):
    """Troca os dados de um gráfico nativo. Os gráficos do modelo apontam para
    uma planilha Excel EXTERNA (vínculo da máquina de quem fez o PPTX): o vínculo
    sai e os dados passam a viver numa planilha embutida no próprio arquivo."""
    from pptx.chart.data import CategoryChartData
    cs = shape.chart._chartSpace
    ext = cs.find('{http://schemas.openxmlformats.org/drawingml/2006/chart}externalData')
    if ext is not None:
        rid = ext.get('{http://schemas.openxmlformats.org/officeDocument/2006/relationships}id')
        cs.remove(ext)
        if rid:
            shape.chart.part.drop_rel(rid) if hasattr(shape.chart.part, 'drop_rel') else None
            rels = shape.chart.part.rels
            if rid in rels:
                rels.pop(rid)
    fc = re.search(r'<c:formatCode>([^<]*)</c:formatCode>', shape.chart.part.blob.decode('utf-8', 'ignore'))
    cd = CategoryChartData(number_format=fc.group(1) if fc else '0.00%')
    cd.categories = categorias
    cd.add_series(shape.chart.plots[0].series[0].name or '', valores)
    shape.chart.replace_data(cd)


def _troca_imagem(shape, png):
    rid = shape._element.blipFill.find(A + 'blip').get(
        '{http://schemas.openxmlformats.org/officeDocument/2006/relationships}embed')
    shape.part.related_part(rid)._blob = png


# -------------------------------------------- página nova: rentabilidade (HYD60)
def _pagina_rentabilidade(prs, depois_de, d, larg_pt=540):
    from pptx.chart.data import CategoryChartData
    from pptx.dml.color import RGBColor
    from pptx.enum.chart import XL_CHART_TYPE, XL_LEGEND_POSITION
    from pptx.util import Pt

    ref = prs.slides[depois_de - 1]
    nova = prs.slides.add_slide(ref.slide_layout)
    for ph in list(nova.placeholders):
        ph._element.getparent().remove(ph._element)

    titulos = [s for s in ref.shapes if s.shape_type == 6
               and any(x.has_text_frame and 'Alocação' in x.text_frame.text for x in s.shapes)]
    uso = [s for s in ref.shapes if s.has_text_frame and 'USO INTERNO' in s.text_frame.text.upper()]

    def titulo(texto, top):
        if not titulos:
            return
        el = copy.deepcopy(titulos[0]._element)
        nova.shapes._spTree.append(el)
        g = nova.shapes[-1]
        dy = Pt(top) - g.top
        g.top = g.top + dy
        for x in g.shapes:
            if x.has_text_frame:
                _troca_runs(x.text_frame.paragraphs[0], [texto])

    for u in uso:
        nova.shapes._spTree.append(copy.deepcopy(u._element))

    VERDE, CINZA, PRETO = RGBColor(0x33, 0x7F, 0x51), RGBColor(0x80, 0x80, 0x80), RGBColor(0, 0, 0)
    titulo('Rentabilidade', 86)
    linhas = d['rent']
    tb = nova.shapes.add_table(len(linhas) + 1, 3, Pt(36), Pt(122), Pt(468), Pt(22 * (len(linhas) + 1))).table
    tb.columns[0].width = Pt(168)
    tb.columns[1].width = Pt(150)
    tb.columns[2].width = Pt(150)
    cab = ['', 'Mês', 'Desde o Início']
    for j, t in enumerate(cab):
        _celula(tb.cell(0, j), t, 9, True, PRETO, j > 0)
    for i, (rot, m, ini) in enumerate(linhas, 1):
        cor = VERDE if i == 1 else CINZA
        _celula(tb.cell(i, 0), rot, 9, True, PRETO, False)
        _celula(tb.cell(i, 1), m, 9, i == 1, cor, True)
        _celula(tb.cell(i, 2), ini, 9, i == 1, cor, True)

    titulo('Rentabilidade Histórica', 300)
    h = d['hist']
    if h:
        cd = CategoryChartData(number_format='0.00%')
        cd.categories = [x.replace('/', '-') for x in h['l']]
        cd.add_series('Fundo', [v / 100 for v in h['f']])
        cd.add_series('CDI', [v / 100 for v in h['c']])
        gf = nova.shapes.add_chart(XL_CHART_TYPE.LINE, Pt(36), Pt(336), Pt(468), Pt(250), cd)
        ch = gf.chart
        ch.has_legend = True
        ch.legend.position = XL_LEGEND_POSITION.BOTTOM
        ch.legend.include_in_layout = False
        ch.legend.font.size = Pt(8)
        ch.font.size = Pt(8)
        ch.font.name = 'Versos'
        ch.font.color.rgb = RGBColor(0x59, 0x59, 0x59)
        for se, cor, w in zip(ch.plots[0].series, (RGBColor(0x4B, 0x9D, 0x74), RGBColor(0x92, 0xD3, 0xAB)), (2.25, 1.75)):
            se.format.line.color.rgb = cor
            se.format.line.width = Pt(w)
            se.smooth = False
            pt = se.points[len(h['l']) - 1]
            pt.data_label.has_text_frame = False
            pt.data_label.show_value = True
            pt.data_label.font.size = Pt(8)
            pt.data_label.font.bold = True
            pt.data_label.font.color.rgb = cor
            from pptx.enum.chart import XL_LABEL_POSITION
            pt.data_label.position = XL_LABEL_POSITION.RIGHT
        va = ch.value_axis
        va.has_major_gridlines = True
        va.major_gridlines.format.line.color.rgb = RGBColor(0xD9, 0xD9, 0xD9)
        va.tick_labels.number_format = '0.00%'
        va.tick_labels.number_format_is_linked = False
        va.format.line.fill.background()
        ca = ch.category_axis
        ca.format.line.color.rgb = RGBColor(0xD9, 0xD9, 0xD9)
        ca.tick_labels.font.size = Pt(7)

    # leva a página nova para logo depois da de referência
    lista = prs.slides._sldIdLst
    el = list(lista)[-1]
    lista.remove(el)
    lista.insert(depois_de, el)


def _celula(c, texto, tam, negrito, cor, centro):
    from pptx.enum.text import PP_ALIGN
    from pptx.util import Pt
    c.fill.background()
    c.margin_left = c.margin_right = Pt(2)
    c.margin_top = c.margin_bottom = Pt(3)
    p = c.text_frame.paragraphs[0]
    p.text = ''
    r = p.add_run()
    r.text = texto
    r.font.size = Pt(tam)
    r.font.bold = negrito
    r.font.color.rgb = cor
    r.font.name = 'Versos'
    p.alignment = PP_ALIGN.CENTER if centro else PP_ALIGN.LEFT
    # linha cinza embaixo de cada célula, como na tabela do e-mail
    tcPr = c._tc.get_or_add_tcPr()
    for lado in ('a:lnL', 'a:lnR', 'a:lnT', 'a:lnB'):
        tag = A + lado[2:]
        velho = tcPr.find(tag)
        if velho is not None:
            tcPr.remove(velho)
    from lxml import etree
    for lado in ('lnL', 'lnR', 'lnT'):
        ln = etree.SubElement(tcPr, A + lado, w='0')
        etree.SubElement(ln, A + 'noFill')
    ln = etree.SubElement(tcPr, A + 'lnB', w='9525')
    sf = etree.SubElement(ln, A + 'solidFill')
    etree.SubElement(sf, A + 'srgbClr', val='BFBFBF')


# ------------------------------------------------------ imagem de setores
def _png_setorial(d, larg_pt, alt_pt, css, html_bloco, navegador):
    """A imagem de setores do Informativo, no tamanho exato da moldura dele."""
    # 1 pt da moldura = 1,65 px: a letra de 11 px do bloco sai com ~7 pt no
    # slide, o tamanho da imagem original
    escala = 4
    w = round(larg_pt * 1.65)
    h = round(alt_pt * 1.65)
    pg = navegador.new_page(viewport={'width': w, 'height': h}, device_scale_factor=escala)
    try:
        pg.set_content('<!doctype html><html><head><meta charset="utf-8"><style>' + css +
                       'html,body{margin:0;background:transparent}'
                       f'.bloco{{width:{w}px;background:transparent;height:{h}px;padding:6px 8px 18px;display:flex}}'
                       '.bloco>div{width:100%;height:100%}'
                       '.b-set{justify-content:space-between;gap:0}'
                       '</style></head><body><div class="bloco">' + html_bloco + '</div></body></html>')
        pg.evaluate('document.fonts && document.fonts.ready')
        pg.wait_for_timeout(250)
        return pg.locator('.bloco').screenshot(omit_background=True)
    finally:
        pg.close()


FONTES_PPT = ('.ttf', '.otf')


def versos_instalada():
    """A Versos está instalada no Windows (para todos ou só para o usuário)?"""
    if os.name != 'nt':
        return False
    pastas = [os.path.join(os.environ.get('WINDIR', r'C:\Windows'), 'Fonts'),
              os.path.join(os.environ.get('LOCALAPPDATA', ''), 'Microsoft', 'Windows', 'Fonts')]
    for pasta in pastas:
        if os.path.isdir(pasta) and any(n.lower().startswith('versos') for n in os.listdir(pasta)):
            return True
    return False


def _carregar_fontes(caminhos):
    """Carrega as fontes na sessão do Windows e avisa os programas abertos.

    Fonte instalada depois do login só aparece para o PowerPoint depois disto
    (ou de sair e entrar de novo no Windows).
    """
    try:
        import ctypes
        gdi = ctypes.windll.gdi32
        for c in caminhos:
            gdi.AddFontResourceW(c)
        HWND_BROADCAST, WM_FONTCHANGE = 0xFFFF, 0x001D
        ctypes.windll.user32.SendMessageTimeoutW(HWND_BROADCAST, WM_FONTCHANGE, 0, 0, 0x0002, 1000, None)
    except Exception:
        pass


def instalar_versos(log):
    """Instala a Versos de assets/fontes/ só para o usuário (sem administrador).

    O PowerPoint só desenha e embute a fonte que está instalada no Windows. Sem
    ela, o PDF do Informativo sai em Calibri/Aptos. O PowerPoint lê .ttf/.otf;
    .woff2 (o formato da web) não serve aqui.
    """
    if os.name != 'nt':
        return False
    pasta_usuario = os.path.join(os.environ.get('LOCALAPPDATA', ''), 'Microsoft', 'Windows', 'Fonts')
    sistema = os.path.join(os.environ.get('WINDIR', r'C:\Windows'), 'Fonts')
    origem = os.path.join(RAIZ, 'assets', 'fontes')
    todos = [n for n in (os.listdir(origem) if os.path.isdir(origem) else [])
             if n.lower().startswith('versos') and n.lower().endswith(FONTES_PPT)]
    # "Versos Light" é a ExtraLight com outro nome (o CDN da marca não tem a
    # Light; o site usa a ExtraLight nos pesos 200 e 300): os modelos pedem as duas
    arquivos = [n for n in todos if not os.path.exists(os.path.join(pasta_usuario, n))
                and not os.path.exists(os.path.join(sistema, n))]
    if todos and not arquivos:
        _carregar_fontes([os.path.join(pasta_usuario, n) for n in todos
                          if os.path.exists(os.path.join(pasta_usuario, n))])
        return True
    if not arquivos:
        log.aviso('informativo', 'fonte Versos não instalada nesta máquina e sem .ttf/.otf em '
                                 'assets/fontes: o PDF do Informativo sai com Calibri/Aptos')
        return False
    import shutil
    import winreg
    destino = os.path.join(os.environ['LOCALAPPDATA'], 'Microsoft', 'Windows', 'Fonts')
    os.makedirs(destino, exist_ok=True)
    chave = winreg.CreateKey(winreg.HKEY_CURRENT_USER,
                             r'Software\Microsoft\Windows NT\CurrentVersion\Fonts')
    for n in arquivos:
        alvo = os.path.join(destino, n)
        shutil.copy2(os.path.join(origem, n), alvo)
        nome = os.path.splitext(n)[0].replace('-', ' ')
        tipo = 'TrueType' if n.lower().endswith('.ttf') else 'OpenType'
        winreg.SetValueEx(chave, f'{nome} ({tipo})', 0, winreg.REG_SZ, alvo)
    winreg.CloseKey(chave)
    _carregar_fontes([os.path.join(destino, n) for n in arquivos])
    log.info(f'  fonte Versos instalada para o usuário ({len(arquivos)} arquivos)')
    return True


def _pdf_powerpoint(pptx, pdf, log):
    """PPTX -> PDF pelo PowerPoint (COM), e o PPTX regravado com a fonte embutida
    (quem abre sem a Versos instalada vê a fonte certa). Sem PowerPoint: False."""
    if os.name != 'nt':
        return False
    embute = versos_instalada()
    script = (
        "$ErrorActionPreference='Stop';"
        "$pp=New-Object -ComObject PowerPoint.Application;"
        f"$p=$pp.Presentations.Open('{pptx}',$false,$false,$false);"
        + (f"$p.SaveAs('{pptx}',24,-1);" if embute else "")
        + f"$p.SaveAs('{pdf}',32);$p.Close();$pp.Quit()")
    try:
        r = subprocess.run(['powershell', '-NoProfile', '-NonInteractive', '-Command', script],
                           capture_output=True, text=True, timeout=240)
        return r.returncode == 0 and os.path.exists(pdf)
    except Exception as e:
        log.aviso('informativo', f'PDF pelo PowerPoint falhou: {e!r}')
        return False


def exportar_informativos(fundos, destino, rend_email, log):
    """fundos: [(dados, pg)] de RenderizadorEmailPrevidencia.fundos()."""
    from pptx import Presentation
    from playwright.sync_api import sync_playwright
    os.makedirs(destino, exist_ok=True)
    instalar_versos(log)
    css = rend_email.r.css_marca + open(os.path.join(RAIZ, 'templates', 'previdencia', 'blocos.css'),
                                         encoding='utf-8').read()
    macros = rend_email.r.env.get_template('previdencia/blocos.html').module
    n = 0
    with sync_playwright() as pw:
        op = {'args': ['--no-sandbox', '--font-render-hinting=none']}
        exe = achar_chromium()
        if exe:
            op['executable_path'] = exe
        nav = pw.chromium.launch(**op)
        try:
            for d, pg in fundos:
                k = d['chave']
                modelo = pg.get('informativo')
                mapa = MAPAS.get(k)
                if not modelo or not mapa:
                    continue
                modelo = modelo if os.path.isabs(modelo) else os.path.join(RAIZ, modelo)
                if not os.path.exists(modelo):
                    log.aviso(k, f'Informativo: modelo não encontrado ({modelo})')
                    continue
                prs = Presentation(modelo)
                S = lambda i: prs.slides[i - 1]
                _mes_e_data(prs, d)
                if 'texto' in mapa:
                    for sh in _acha(S(mapa['texto'][0]), mapa['texto'][1]):
                        _escreve_texto(sh, d['texto_md'])
                if 'pl' in mapa:
                    for sh in _acha(S(mapa['pl'][0]), mapa['pl'][1]):
                        _pl(sh, d)
                if 'duration' in mapa:
                    for sh in _acha(S(mapa['duration'][0]), mapa['duration'][1]):
                        _duration(sh, d)
                _alocacao(S(mapa['alocacao'][0]), mapa['alocacao'][1], d)
                for sh in _acha(S(mapa['rating'][0]), mapa['rating'][1], lambda s: s.has_chart):
                    _grafico(sh, [r[0] for r in d['rating']], [r[2] / 100 for r in d['rating']])
                if 'composicao' in mapa:
                    comp = list(reversed(d['composicao']))      # barras: de baixo para cima
                    for sh in _acha(S(mapa['composicao'][0]), mapa['composicao'][1], lambda s: s.has_chart):
                        _grafico(sh, [c[0] for c in comp], [c[2] / 100 for c in comp])
                for sh in _acha(S(mapa['setorial'][0]), mapa['setorial'][1]):
                    png = _png_setorial(d, sh.width / 12700, sh.height / 12700, css,
                                        str(macros.setorial(d)), nav)
                    _troca_imagem(sh, png)
                if mapa.get('pagina_rentabilidade'):
                    _pagina_rentabilidade(prs, mapa['pagina_rentabilidade'], d)

                base = f'Informativo - {d["nome_curto"].replace("TIVIO ", "")} - {d["mes"]} {d["ano"]}'
                pptx = os.path.join(destino, base + '.pptx')
                prs.save(pptx)
                log.gerado(pptx, 'pptx')
                pdf = os.path.join(destino, base + '.pdf')
                if _pdf_powerpoint(os.path.abspath(pptx), os.path.abspath(pdf), log):
                    log.gerado(pdf, 'pdf')
                else:
                    log.aviso(k, 'Informativo em PDF não gerado: precisa do PowerPoint instalado '
                                 '(o PPTX saiu; dá para salvar como PDF por ele)')
                n += 1
        finally:
            nav.close()
    return n


# ------------------------------------------------------------------- e-mail
def _eml(html, pngs, assunto):
    """E-mail com as imagens embutidas (cid:), que o Outlook abre como rascunho."""
    from email.mime.image import MIMEImage
    from email.mime.multipart import MIMEMultipart
    from email.mime.text import MIMEText
    corpo = html
    for nome in pngs:
        corpo = corpo.replace(f'src="{nome}"', f'src="cid:{nome}"')
    raiz = MIMEMultipart('related')
    raiz['Subject'] = assunto
    raiz['X-Unsent'] = '1'
    raiz.attach(MIMEText(corpo, 'html', 'utf-8'))
    for nome, dados in pngs.items():
        img = MIMEImage(dados, 'png')
        img.add_header('Content-ID', f'<{nome}>')
        img.add_header('Content-Disposition', 'inline', filename=nome)
        raiz.attach(img)
    return raiz.as_bytes(), corpo


def _oft(html_cid, pasta, pngs, assunto, destino, log):
    """Modelo do Outlook (.oft), pelo próprio Outlook. Sem Outlook: False."""
    if os.name != 'nt':
        return False
    anexos = ''.join(
        f"$a=$m.Attachments.Add('{os.path.join(pasta, n)}');"
        f"$a.PropertyAccessor.SetProperty('http://schemas.microsoft.com/mapi/proptag/0x3712001F','{n}');"
        for n in pngs)
    corpo = os.path.join(pasta, '_corpo_oft.html')
    with open(corpo, 'w', encoding='utf-8') as f:
        f.write(html_cid)
    script = (
        "$ErrorActionPreference='Stop';"
        "$o=New-Object -ComObject Outlook.Application;$m=$o.CreateItem(0);"
        f"$m.Subject='{assunto}';{anexos}"
        f"$m.HTMLBody=[IO.File]::ReadAllText('{corpo}',[Text.Encoding]::UTF8);"
        f"$m.SaveAs('{destino}',2)")
    try:
        r = subprocess.run(['powershell', '-NoProfile', '-NonInteractive', '-Command', script],
                           capture_output=True, text=True, timeout=60)
        ok = r.returncode == 0 and os.path.exists(destino)
        if not ok and r.stderr:
            log.aviso('e-mail previdência', 'modelo .oft não gerado: ' + r.stderr.strip()[:200])
        return ok
    except Exception as e:
        log.aviso('e-mail previdência', f'modelo .oft não gerado: {e!r}')
        return False
    finally:
        if os.path.exists(corpo):
            os.remove(corpo)


def exportar_email(pasta_central, destino, rend_email, fundos, log, oft=False):
    from playwright.sync_api import sync_playwright
    origem = os.path.join(pasta_central, rend_email.ARQUIVO)
    if not os.path.exists(origem):
        return 0
    total = 0
    with sync_playwright() as pw:
        op = {'args': ['--no-sandbox', '--font-render-hinting=none']}
        exe = achar_chromium()
        if exe:
            op['executable_path'] = exe
        nav = pw.chromium.launch(**op)
        try:
            pg = nav.new_page(viewport={'width': 1700, 'height': 1000}, device_scale_factor=3)
            pg.goto('file:///' + os.path.abspath(origem).replace(os.sep, '/'), wait_until='load')
            pg.wait_for_function('typeof setFundo === "function"', timeout=20000)
            for d, _ in fundos:
                k = d['chave']
                pg.evaluate('k => setFundo(k)', k)
                pg.evaluate('document.body.classList.add("exportando")')
                pg.evaluate('document.fonts && document.fonts.ready')
                pg.wait_for_timeout(900)
                pasta = os.path.join(destino, k)
                os.makedirs(pasta, exist_ok=True)
                for velho in os.listdir(pasta):
                    if velho.lower().endswith(('.png', '.html', '.eml', '.oft')):
                        os.remove(os.path.join(pasta, velho))
                arquivos = rend_email.arquivos(d)
                pngs = {}
                for b, nome in arquivos.items():
                    el = pg.locator(f'.fundo[data-fundo="{k}"] .bloco[data-bloco="{b}"]')
                    pngs[nome] = el.screenshot()
                    with open(os.path.join(pasta, nome), 'wb') as f:
                        f.write(pngs[nome])
                html = rend_email.disparo(d, arquivos)
                with open(os.path.join(pasta, f'email-{k}.html'), 'w', encoding='utf-8') as f:
                    f.write(html)
                assunto = f'Saiba mais sobre o fundo {d["nome_curto"]}'
                eml, html_cid = _eml(html, pngs, assunto)
                with open(os.path.join(pasta, f'email-{k}.eml'), 'wb') as f:
                    f.write(eml)
                if oft:
                    _oft(html_cid, os.path.abspath(pasta), pngs, assunto,
                         os.path.abspath(os.path.join(pasta, f'email-{k}.oft')), log)
                total += 1
                log.gerado(pasta, 'png')
        finally:
            nav.close()
    return total
