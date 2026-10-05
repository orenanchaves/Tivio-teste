# -*- coding: utf-8 -*-
"""PPTX fiel à folha HTML, com o texto editável.

A abordagem é a da skill ppt-master (github.com/hugohe3/ppt-master), rota
"Image to PPTX": cada slide é a página reconstruída em camadas registradas —

  1. fundo: a folha renderizada no Chromium SEM o texto (cores, faixas,
     barras, treemap, gráficos, logos, selos), numa imagem em 2x;
  2. por cima: cada bloco de texto da folha como caixa de texto NATIVA do
     PowerPoint, na posição, largura, corpo, peso, cor e alinhamento medidos
     no próprio navegador.

O slide fica idêntico ao PDF e o texto continua editável no PowerPoint — o que
o PPTX "nativo" anterior não conseguia: ele redesenhava a folha com 10 a 18
formas e saía com outro layout.

Rótulos dentro de SVG (eixos do gráfico, logos) ficam na imagem: são parte do
desenho, não texto de edição.
"""
import io
import os
import re

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Emu, Pt

# A4 retrato, igual à folha (1000 x 1414 px)
LARG_EMU = 7560000
ALT_EMU = 10692000
FOLHA_PX = 1000

# Coleta, no navegador, os blocos de texto de cada folha: o elemento que tem
# texto direto e cujo pai não tem (o <b> dentro do <p> vira run do <p>).
COLETA = r"""
() => {
  const folhas = [...document.querySelectorAll('.rcard')];
  const cor = c => {
    const m = (c || '').match(/rgba?\(([^)]+)\)/);
    if (!m) return null;
    const p = m[1].split(',').map(s => parseFloat(s));
    if (p.length > 3 && p[3] < 0.05) return null;
    return [p[0], p[1], p[2]];
  };
  const temTexto = el => [...el.childNodes].some(n => n.nodeType === 3 && n.nodeValue.trim());
  const estilo = el => {
    const s = getComputedStyle(el);
    return { cor: cor(s.color), peso: parseInt(s.fontWeight) || 400,
             italico: s.fontStyle === 'italic', corpo: parseFloat(s.fontSize),
             caixa: s.textTransform };
  };
  return folhas.map(folha => {
    const base = folha.getBoundingClientRect();
    const k = 1000 / base.width;
    const blocos = [];
    folha.querySelectorAll('*').forEach(el => {
      if (el.closest('svg') || !temTexto(el)) return;
      // dentro de um bloco que já tem texto: vira run dele, não caixa à parte
      // (num flex o <b> vira bloco, e saía duplicado)
      for (let a = el.parentElement; a && a !== folha; a = a.parentElement)
        if (temTexto(a)) return;
      const s = getComputedStyle(el);
      if (s.visibility === 'hidden' || s.display === 'none' || parseFloat(s.opacity) === 0) return;
      const r = el.getBoundingClientRect();
      if (r.width < 1 || r.height < 1) return;
      const pl = parseFloat(s.paddingLeft), pr = parseFloat(s.paddingRight);
      const pt = parseFloat(s.paddingTop), pb = parseFloat(s.paddingBottom);
      // a caixa do texto: do primeiro caractere ao fim, medida pelo texto de
      // verdade (não pela célula, que é mais alta; nem pelo quadradinho da legenda)
      const andar = document.createTreeWalker(el, NodeFilter.SHOW_TEXT);
      let primeiro = null, nn;
      while ((nn = andar.nextNode())) { if (nn.nodeValue.trim()) { primeiro = nn; break; } }
      const faixa = document.createRange();
      faixa.selectNodeContents(el);
      if (primeiro) faixa.setStart(primeiro, primeiro.nodeValue.search(/\S/));
      const t = faixa.getBoundingClientRect();
      const runs = [];
      el.childNodes.forEach(n => {
        if (n.nodeType === 3) {
          if (n.nodeValue) runs.push(Object.assign({ t: n.nodeValue }, estilo(el)));
        } else if (n.nodeType === 1) {
          if (n.tagName === 'BR') { runs.push({ t: '\n' }); return; }
          if (n.closest('svg') || n.tagName === 'svg') return;
          const txt = n.innerText;
          if (txt) runs.push(Object.assign({ t: txt }, estilo(n)));
        }
      });
      const caixa = el.innerText.trim();
      if (!caixa) return;
      const lh = parseFloat(s.lineHeight) || parseFloat(s.fontSize) * 1.2;
      // na horizontal, a área de conteúdo (o alinhamento do CSS decide onde o
      // texto cai); alinhado à esquerda, começa onde o texto começa
      const esq = (s.textAlign === 'left' || s.textAlign === 'start' || s.textAlign === 'justify')
        ? Math.max(r.left + pl, t.left) : r.left + pl;
      const dir = r.right - pr;
      blocos.push({
        x: (esq - base.left) * k, y: (t.top - base.top) * k,
        w: Math.max(t.width, dir - esq) * k, h: t.height * k,
        linhas: Math.round(t.height / lh),
        alinha: s.textAlign, lh: lh * k, caixaAlta: s.textTransform === 'uppercase',
        runs: runs, texto: caixa,
      });
    });
    return blocos;
  });
}
"""

# Esconde o texto HTML (vai como caixa nativa por cima). SVG fica de fora: o
# logo usa currentColor, e a cor de cada SVG é fixada antes (FIXA_SVG).
ESCONDE_TEXTO = """
.rcard, .rcard :not(svg):not(svg *) { color: transparent !important;
  -webkit-text-fill-color: transparent !important; text-shadow: none !important; }
"""
FIXA_SVG = """() => document.querySelectorAll('.rcard svg').forEach(s =>
  s.style.setProperty('color', getComputedStyle(s).color, 'important'))"""


def _rgb(c):
    return RGBColor(*[max(0, min(255, int(round(v)))) for v in c]) if c else None


ALINHA = {'center': PP_ALIGN.CENTER, 'right': PP_ALIGN.RIGHT, 'end': PP_ALIGN.RIGHT,
          'justify': PP_ALIGN.JUSTIFY}


class ExportadorPPTXFiel:
    def __init__(self, log, fonte='Versos'):
        self.log = log
        self.fonte = fonte

    def exportar(self, browser, origem, destino):
        """`browser` é o Chromium já aberto pelo ExportadorPDF."""
        ctx = browser.new_context(viewport={'width': 1000, 'height': 1414},
                                  device_scale_factor=2)
        pg = ctx.new_page()
        try:
            pg.goto('file://' + os.path.abspath(origem), wait_until='load')
            # zoom 1: a folha em 1000x1414 exatos (a tela encolhe por media query)
            pg.evaluate("document.body.classList.add('tv-exportando')")
            pg.wait_for_function('document.fonts.ready.then(()=>true)', timeout=15000)
            pg.evaluate('window.tvAjustarTextos && window.tvAjustarTextos()')
            pg.wait_for_timeout(400)
            blocos = pg.evaluate(COLETA)
            pg.evaluate(FIXA_SVG)
            pg.add_style_tag(content=ESCONDE_TEXTO)
            pg.wait_for_timeout(150)
            fundos = [f.screenshot(type='png') for f in pg.query_selector_all('.rcard')]
        finally:
            ctx.close()

        prs = Presentation()
        prs.slide_width, prs.slide_height = Emu(LARG_EMU), Emu(ALT_EMU)
        k = LARG_EMU / FOLHA_PX                      # EMU por px da folha
        vazio = prs.slide_layouts[6]
        n_txt = 0
        for fundo, caixas in zip(fundos, blocos):
            sl = prs.slides.add_slide(vazio)
            sl.shapes.add_picture(io.BytesIO(fundo), 0, 0, Emu(LARG_EMU), Emu(ALT_EMU))
            for b in caixas:
                self._caixa(sl, b, k)
                n_txt += 1
        os.makedirs(os.path.dirname(destino) or '.', exist_ok=True)
        prs.save(destino)
        return n_txt

    def _caixa(self, slide, b, k):
        folga = 1.04 if b['linhas'] <= 1 else 1.0
        tb = slide.shapes.add_textbox(Emu(int(b['x'] * k)), Emu(int(b['y'] * k)),
                                      Emu(int(b['w'] * k * folga)), Emu(int(max(b['h'], 4) * k)))
        tf = tb.text_frame
        tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
        # uma linha só não pode quebrar no PowerPoint por diferença de métrica
        tf.word_wrap = b['linhas'] > 1
        tf.vertical_anchor = MSO_ANCHOR.TOP
        tf.auto_size = None
        par = tf.paragraphs[0]
        par.alignment = ALINHA.get(b['alinha'], PP_ALIGN.LEFT)
        # entrelinha em pontos: px da folha -> pt do slide (A4 = 595,3 pt de largura)
        par.line_spacing = Pt(b['lh'] * 595.3 / FOLHA_PX)
        for r in b['runs']:
            texto = r.get('t', '')
            if not texto:
                continue
            if texto == '\n':
                par = tf.add_paragraph()
                par.alignment = ALINHA.get(b['alinha'], PP_ALIGN.LEFT)
                par.line_spacing = Pt(b['lh'] * 595.3 / FOLHA_PX)
                continue
            texto = re.sub(r'\s+', ' ', texto)
            if r.get('caixa') == 'uppercase' or b.get('caixaAlta'):
                texto = texto.upper()
            run = par.add_run()
            run.text = texto
            f = run.font
            f.name = self.fonte
            f.size = Pt((r.get('corpo') or 12) * 595.3 / FOLHA_PX)
            f.bold = (r.get('peso') or 400) >= 600
            f.italic = bool(r.get('italico'))
            c = _rgb(r.get('cor'))
            if c is not None:
                f.color.rgb = c
        # tira o espaço da ponta (o HTML ignora; o PowerPoint desenha)
        for p in tf.paragraphs:
            if p.runs:
                p.runs[0].text = p.runs[0].text.lstrip()
                p.runs[-1].text = p.runs[-1].text.rstrip()
