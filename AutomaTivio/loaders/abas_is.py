# -*- coding: utf-8 -*-
"""Abas de Investment Solutions da dados_mensais.xlsx ("Low Vol", "TOP Credit").

Cada aba é a folha de trabalho de um fundo de previdência da Itaú, montada pela
gestão. Os decks desses fundos já nasceram ligados a ela: os gráficos do PPTX
guardam a fórmula de onde vieram ('Low Vol'!$T$6:$T$13). A automação lê
exatamente esses intervalos, então o gráfico mostra o mesmo recorte que a
gestão escolheu, sem mapa de colunas para manter aqui.

O que não é gráfico sai pelo rótulo, não pela posição:
  - "Vol" na linha do nome do fundo        -> volatilidade 12 meses
  - "PnL (bps)"                            -> composição da rentabilidade
    (cabeçalho "Estratégia | Alpha | Bench | Total" e as linhas até "Total")
"""
import re

import openpyxl

_cache = {}


def ler_abas(caminho, nomes):
    """{aba: [[valores]]} das abas pedidas que existem na planilha."""
    chave = (caminho, tuple(sorted(nomes)))
    if chave in _cache:
        return _cache[chave]
    wb = openpyxl.load_workbook(caminho, read_only=True, data_only=True)
    out = {}
    for nome in nomes:
        if nome in wb.sheetnames:
            out[nome] = [list(r) for r in wb[nome].iter_rows(values_only=True)]
    wb.close()
    _cache[chave] = out
    return out


def _col(letras):
    n = 0
    for ch in letras.upper():
        n = n * 26 + (ord(ch) - 64)
    return n - 1


_REF = re.compile(r"^(?:'([^']+)'|([^'!]+))!\$?([A-Z]+)\$?(\d+)(?::\$?([A-Z]+)\$?(\d+))?$")


def faixa(abas, formula):
    """Valores do intervalo da fórmula ('Low Vol'!$T$6:$T$13), ou None se a aba
    não está entre as lidas (ex.: gráfico ligado a outra planilha)."""
    m = _REF.match((formula or '').strip())
    if not m:
        return None
    aba = m.group(1) or m.group(2)
    if aba not in abas:
        return None
    g = abas[aba]
    c0, r0 = _col(m.group(3)), int(m.group(4)) - 1
    c1 = _col(m.group(5)) if m.group(5) else c0
    r1 = int(m.group(6)) - 1 if m.group(6) else r0
    out = []
    for r in range(r0, r1 + 1):
        for c in range(c0, c1 + 1):
            out.append(g[r][c] if r < len(g) and c < len(g[r]) else None)
    return out


def _acha(grade, texto, max_linhas=60):
    alvo = texto.strip().lower()
    for r, linha in enumerate(grade[:max_linhas]):
        for c, v in enumerate(linha):
            if isinstance(v, str) and v.strip().lower() == alvo:
                return r, c
    return None


def volatilidade(grade):
    """Valor à direita do rótulo "Vol" (fração: 0,0093 = 0,93%)."""
    pos = _acha(grade, 'Vol', 10)
    if not pos:
        return None
    r, c = pos
    v = grade[r][c + 1] if c + 1 < len(grade[r]) else None
    return float(v) if isinstance(v, (int, float)) else None


def pnl(grade):
    """[(estratégia, alpha, bench, total)] do bloco "PnL (bps)", até a linha Total."""
    pos = _acha(grade, 'PnL (bps)')
    if not pos:
        return []
    r, c = pos
    out = []
    for linha in grade[r + 2:r + 40]:                  # pula o cabeçalho Estratégia/Alpha/...
        nome = linha[c] if c < len(linha) else None
        if not isinstance(nome, str) or not nome.strip():
            break
        vals = [linha[c + k] if c + k < len(linha) else None for k in (1, 2, 3)]
        out.append((nome.strip(), *[float(v) if isinstance(v, (int, float)) else None for v in vals]))
        if nome.strip().lower() == 'total':
            break
    return out
