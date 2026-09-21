# -*- coding: utf-8 -*-
"""Gera a versao .docx do curriculo (estrutura ATS-safe: coluna unica, sem tabelas/caixas)."""
from docx import Document
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

DOC = Document()

# margens
for s in DOC.sections:
    s.top_margin = s.bottom_margin = Inches(0.5)
    s.left_margin = s.right_margin = Inches(0.45)

RIGHT = Inches(7.1)  # 8.5 - 0.45*2 = 7.6; tab um pouco antes da borda

normal = DOC.styles["Normal"]
normal.font.name = "Calibri"
normal.font.size = Pt(10.5)
normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Calibri")
pf = normal.paragraph_format
pf.space_before = Pt(0); pf.space_after = Pt(0); pf.line_spacing = 1.05


def par(align=None, space_after=0, space_before=0):
    p = DOC.add_paragraph()
    if align is not None:
        p.alignment = align
    p.paragraph_format.space_after = Pt(space_after)
    p.paragraph_format.space_before = Pt(space_before)
    return p


def run(p, texto, bold=False, size=10.5):
    r = p.add_run(texto)
    r.bold = bold
    r.font.size = Pt(size)
    r.font.color.rgb = RGBColor(0, 0, 0)
    return r


def borda_superior(p):
    """Regua cinza acima do paragrafo (equivalente ao filete do PDF)."""
    pPr = p._p.get_or_add_pPr()
    bdr = OxmlElement("w:pBdr")
    top = OxmlElement("w:top")
    top.set(qn("w:val"), "single")
    top.set(qn("w:sz"), "8")
    top.set(qn("w:space"), "6")
    top.set(qn("w:color"), "9A9A9A")
    bdr.append(top)
    pPr.append(bdr)


def secao(titulo):
    p = par(WD_ALIGN_PARAGRAPH.CENTER, space_after=5, space_before=9)
    borda_superior(p)
    run(p, titulo, bold=True, size=11.5)


def cargo(titulo, periodo, empresa):
    p = par(space_before=8)
    p.paragraph_format.tab_stops.add_tab_stop(RIGHT, WD_TAB_ALIGNMENT.RIGHT)
    run(p, titulo, bold=True)
    run(p, "\t" + periodo)
    p2 = par()
    run(p2, empresa, bold=True)


def bullets(itens):
    for i, t in enumerate(itens):
        p = DOC.add_paragraph(style="List Bullet")
        p.paragraph_format.space_after = Pt(1.5)
        p.paragraph_format.space_before = Pt(3 if i == 0 else 0)
        p.paragraph_format.left_indent = Inches(0.3)
        p.paragraph_format.line_spacing = 1.05
        run(p, t)


# ------------------------------------------------------------------ cabecalho
p = par(WD_ALIGN_PARAGRAPH.CENTER)
run(p, "Renata Xavier", bold=True, size=16)
p = par(WD_ALIGN_PARAGRAPH.CENTER, space_after=2)
run(p, "Analista de Comunicação Interna | Endomarketing | Employer Branding | Cultura e Engajamento", bold=True, size=11)
p = par(WD_ALIGN_PARAGRAPH.CENTER)
run(p, "São Paulo, SP • (11) 96476-9808 • renatadiasx@hotmail.com • linkedin.com/in/renatadxavier")

# ------------------------------------------------------------------ resumo
secao("Resumo Profissional")
p = par(WD_ALIGN_PARAGRAPH.JUSTIFY)
run(p, "Profissional de Comunicação Interna e Endomarketing com mais de 6 anos de experiência em empresas de "
       "grande porte, responsável por planos de comunicação, gestão de canais internos, campanhas de cultura "
       "organizacional, employer branding e eventos corporativos. Conduziu a comunicação de uma operação nacional "
       "com mais de 800 colaboradores e 26 unidades no Brasil, liderando a frente de Diversidade, Equidade e Inclusão "
       "(DE&I) e o censo de diversidade da companhia. Experiência em gestão de agências e fornecedores, comunicação "
       "de liderança, produção de conteúdo e fortalecimento da marca empregadora em canais digitais. Formada em "
       "Propaganda e Marketing, com perfil estratégico, colaborativo e orientado a engajamento e clima organizacional.")

# ------------------------------------------------------------------ competencias
secao("Competências-Chave")
COMP = [
    ("Comunicação & Endomarketing: ", "planejamento e execução de plano de comunicação interna, comunicados corporativos, calendário editorial, comunicação de liderança e executiva, comunicação de mudança, storytelling e produção de conteúdo, campanhas institucionais e de engajamento."),
    ("Canais & Marca Empregadora: ", "gestão de canais internos (rede social corporativa, intranet, e-mail marketing, TV corporativa e murais), employer branding, LinkedIn corporativo e redes sociais, presença digital da marca."),
    ("Cultura, Pessoas & DE&I: ", "cultura organizacional, clima e engajamento, Diversidade, Equidade e Inclusão (DE&I), comitês e censo de diversidade, avaliação de desempenho e calibração, integração e onboarding de colaboradores."),
    ("Projetos & Stakeholders: ", "eventos corporativos de pequeno a grande porte, gestão de agências e fornecedores, briefing e acompanhamento de entregas, patrocínios e leis de incentivo (Lei Rouanet e Lei de Incentivo ao Esporte), atendimento a clientes internos."),
    ("Ferramentas: ", "Pacote Office (Word, PowerPoint, Excel), Trello, plataformas de e-mail marketing, redes sociais corporativas, Adobe Photoshop e Illustrator (intermediário)."),
]
for rot, txt in COMP:
    p = par(WD_ALIGN_PARAGRAPH.JUSTIFY, space_after=2)
    run(p, rot, bold=True)
    run(p, txt)

# ------------------------------------------------------------------ experiencia
secao("Experiência Profissional")

cargo("Analista de Recursos Humanos — Comunicação Interna e Endomarketing", "Junho de 2026 – Atual", "Zeentech")
bullets([
    "Estruturo e executo o plano de comunicação interna da companhia, da integração de novos colaboradores ao engajamento contínuo, definindo pauta, canais e calendário editorial.",
    "Produzo e distribuo comunicados corporativos e faço a gestão dos canais internos de comunicação, garantindo padronização de linguagem e alinhamento à identidade da marca.",
    "Gerencio as redes sociais da empresa, com planejamento de conteúdo e produção de posts voltados ao fortalecimento da marca empregadora nos canais digitais.",
    "Conduzo a gestão da agência de marketing: briefing, direcionamento de demandas, acompanhamento de entregas e alinhamento das ações à estratégia da marca.",
    "Desenvolvo campanhas institucionais e de endomarketing integradas a estratégias de employer branding.",
    "Realizo curadoria e gestão de patrocínios via leis de incentivo (Lei Rouanet e Lei de Incentivo ao Esporte), avaliando projetos e conectando as oportunidades aos objetivos de negócio.",
    "Prospecto e viabilizo patrocínios em eventos do setor e de marketing, ampliando a presença institucional da marca.",
    "Organizo eventos corporativos e institucionais de pequeno a grande porte, do briefing à execução.",
])

cargo("Analista de Comunicação Interna e Endomarketing Pleno", "Setembro de 2021 – Abril de 2026", "Total Express")
bullets([
    "Conduzi o plano de comunicação interna de uma operação nacional com mais de 800 colaboradores e 26 unidades no Brasil, articulando agentes de comunicação locais e garantindo padronização e efetividade na disseminação de informações.",
    "Geri o ecossistema de canais internos — rede social corporativa, e-mail, TV corporativa e murais físicos —, definindo formatos, periodicidade e fluxo de comunicados corporativos.",
    "Criei campanhas de comunicação e endomarketing com potencial de divulgação externa, integrando employer branding e fortalecendo a presença da marca no LinkedIn e demais canais digitais.",
    "Liderei a frente de Diversidade e Inclusão: comitê com reuniões mensais, campanhas de conscientização e sensibilização, workshops e palestras corporativas, consolidando uma cultura inclusiva.",
    "Liderei o censo de diversidade de 2025, da coleta à análise dos dados, transformando os resultados em oportunidades de melhoria e estratégias inclusivas para a empresa.",
    "Planejei e executei eventos institucionais de pequeno, médio e grande porte — convenções, palestras, coffee breaks e encontros para executivos.",
    "Conduzi o ciclo de Avaliação de Desempenho, com uso da plataforma de avaliação, suporte a gestores e calibrações de mais de 800 colaboradores.",
    "Estruturei e acompanhei grandes projetos, como lançamentos de programas e iniciativas para colaboradores, assegurando ações estratégicas de endomarketing alinhadas aos objetivos organizacionais.",
    "Gerenciei agências de comunicação, fornecedores e clientes internos, assegurando qualidade, prazo e aderência ao posicionamento da marca.",
])

cargo("Estagiária de Marketing", "Junho de 2020 – Agosto de 2021", "Grupo UAI")
bullets([
    "Respondi pelas atividades de Marketing e Endomarketing da empresa, da definição à execução das estratégias de comunicação.",
    "Planejei e produzi conteúdo para redes sociais: cronograma editorial, posts, vídeos e redação de legendas alinhadas à identidade da marca.",
    "Criei materiais institucionais, apresentações e conteúdos de apoio.",
    "Desenvolvi ações de comunicação, campanhas, pesquisas de satisfação e eventos corporativos.",
    "Apoiei a estruturação e atualização do site institucional, contribuindo para a melhoria da presença digital da empresa.",
])

# ------------------------------------------------------------------ formacao
secao("Formação Acadêmica")
p = par()
p.paragraph_format.tab_stops.add_tab_stop(RIGHT, WD_TAB_ALIGNMENT.RIGHT)
run(p, "Bacharelado em Propaganda e Marketing", bold=True)
run(p, "\tConclusão: Dezembro de 2022")
p = par()
run(p, "Universidade Paulista (UNIP) — Santana de Parnaíba, SP")

# ------------------------------------------------------------------ idiomas
secao("Idiomas")
p = par()
run(p, "Português", bold=True); run(p, " — Nativo")
p = par()
run(p, "Inglês", bold=True); run(p, " — Intermediário (leitura e escrita)")

DOC.save("/home/user/Tivio-teste/curriculo/Renata-Xavier-Comunicacao-Interna-Endomarketing.docx")
print("docx gerado")
