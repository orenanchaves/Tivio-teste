# -*- coding: utf-8 -*-
"""Gera HTML e DOCX de cada versao do curriculo. O PDF sai do HTML via render.js."""
import html as _html
import re
import sys

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

from conteudo import CONTATO, FORMACAO, IDIOMAS, NOME, VERSOES

NEGRITO = re.compile(r"\*\*(.+?)\*\*")


def partes(texto):
    """Quebra '**x** y' em [(trecho, negrito?), ...]."""
    out, pos = [], 0
    for m in NEGRITO.finditer(texto):
        if m.start() > pos:
            out.append((texto[pos:m.start()], False))
        out.append((m.group(1), True))
        pos = m.end()
    if pos < len(texto):
        out.append((texto[pos:], False))
    return out


# ===================================================================== HTML
CSS = """
  @page { size: Letter; margin: 0.5in 0.45in; }
  html { -webkit-print-color-adjust: exact; print-color-adjust: exact; }
  body { font-family: "Carlito","Calibri","Liberation Sans",sans-serif; font-size: 10.5pt;
         line-height: 1.17; color: #000; margin: 0; }
  .nome { font-size: 16pt; font-weight: bold; text-align: center; margin: 0 0 1pt; letter-spacing: .3pt; }
  .headline { font-size: 11pt; font-weight: bold; text-align: center; margin: 0 0 3pt; }
  .contato { font-size: 10.5pt; text-align: center; margin: 0; }
  .secao { font-size: 11.5pt; font-weight: bold; text-align: center; border-top: 1.2pt solid #9a9a9a;
           padding-top: 6pt; margin: 10pt 0 6pt; }
  .resumo { text-align: justify; margin: 0; }
  .comp { margin: 0; }
  .comp div { margin-bottom: 2.5pt; text-align: justify; }
  .comp .rot { font-weight: bold; }
  .cargo { margin-top: 9pt; page-break-after: avoid; break-after: avoid; }
  .cargo:first-of-type { margin-top: 0; }
  .linha { display: flex; justify-content: space-between; align-items: baseline; gap: 12pt; }
  .linha .dir { text-align: right; white-space: nowrap; }
  .titulo, .empresa { font-weight: bold; }
  ul { margin: 4pt 0 0; padding-left: 26pt; }
  li { text-align: justify; margin-bottom: 2pt; padding-left: 4pt; }
  li::marker { font-size: 10pt; }
  .simples div { margin-bottom: 1.5pt; }
"""


def rich_html(texto):
    return "".join(f"<b>{_html.escape(t)}</b>" if b else _html.escape(t) for t, b in partes(texto))


def gerar_html(v):
    p = [f'<!DOCTYPE html>\n<html lang="pt-BR">\n<head>\n<meta charset="utf-8">',
         f'<title>{_html.escape(NOME)} — {_html.escape(v["headline"].split(" | ")[0])}</title>',
         f"<style>{CSS}</style>\n</head>\n<body>"]
    p.append(f'<h1 class="nome">{_html.escape(NOME)}</h1>')
    p.append(f'<p class="headline">{_html.escape(v["headline"])}</p>')
    p.append(f'<p class="contato">{_html.escape(CONTATO)}</p>')

    p.append('<div class="secao">Resumo Profissional</div>')
    p.append(f'<p class="resumo">{rich_html(v["resumo"])}</p>')

    p.append('<div class="secao">Competências-Chave</div><div class="comp">')
    for rot, txt in v["competencias"]:
        p.append(f'<div><span class="rot">{_html.escape(rot)}:</span> {_html.escape(txt)}</div>')
    p.append("</div>")

    p.append('<div class="secao">Experiência Profissional</div>')
    for titulo, periodo, empresa, bs in v["experiencias"]:
        p.append('<div class="cargo"><div class="linha">'
                 f'<span class="titulo">{_html.escape(titulo)}</span>'
                 f'<span class="dir">{_html.escape(periodo)}</span></div>'
                 f'<div class="empresa">{_html.escape(empresa)}</div></div><ul>')
        p += [f"<li>{rich_html(b)}</li>" for b in bs]
        p.append("</ul>")

    p.append('<div class="secao">Formação Acadêmica</div>')
    p.append('<div class="cargo" style="margin-top:0"><div class="linha">'
             f'<span class="titulo">{_html.escape(FORMACAO[0])}</span>'
             f'<span class="dir">{_html.escape(FORMACAO[1])}</span></div>'
             f'<div>{_html.escape(FORMACAO[2])}</div></div>')

    p.append('<div class="secao">Idiomas</div><div class="simples">')
    p += [f"<div><b>{_html.escape(i)}</b> — {_html.escape(n)}</div>" for i, n in IDIOMAS]
    p.append("</div>\n</body>\n</html>")
    open(v["html"], "w").write("\n".join(p))


# ===================================================================== DOCX
RIGHT = Inches(7.1)


def gerar_docx(v):
    doc = Document()
    for s in doc.sections:
        s.top_margin = s.bottom_margin = Inches(0.5)
        s.left_margin = s.right_margin = Inches(0.45)

    normal = doc.styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(10.5)
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Calibri")
    pf = normal.paragraph_format
    pf.space_before = pf.space_after = Pt(0)
    pf.line_spacing = 1.05

    def par(align=None, after=0, before=0):
        p = doc.add_paragraph()
        if align is not None:
            p.alignment = align
        p.paragraph_format.space_after = Pt(after)
        p.paragraph_format.space_before = Pt(before)
        return p

    def run(p, texto, bold=False, size=10.5):
        r = p.add_run(texto)
        r.bold = bold
        r.font.size = Pt(size)
        r.font.color.rgb = RGBColor(0, 0, 0)

    def rich(p, texto):
        for t, b in partes(texto):
            run(p, t, bold=b)

    def secao(titulo):
        p = par(WD_ALIGN_PARAGRAPH.CENTER, after=5, before=9)
        pPr = p._p.get_or_add_pPr()
        bdr, top = OxmlElement("w:pBdr"), OxmlElement("w:top")
        for k, val in (("val", "single"), ("sz", "8"), ("space", "6"), ("color", "9A9A9A")):
            top.set(qn(f"w:{k}"), val)
        bdr.append(top)
        pPr.append(bdr)
        run(p, titulo, bold=True, size=11.5)

    # cabecalho
    p = par(WD_ALIGN_PARAGRAPH.CENTER); run(p, NOME, bold=True, size=16)
    p = par(WD_ALIGN_PARAGRAPH.CENTER, after=2); run(p, v["headline"], bold=True, size=11)
    p = par(WD_ALIGN_PARAGRAPH.CENTER); run(p, CONTATO)

    secao("Resumo Profissional")
    rich(par(WD_ALIGN_PARAGRAPH.JUSTIFY), v["resumo"])

    secao("Competências-Chave")
    for rot, txt in v["competencias"]:
        p = par(WD_ALIGN_PARAGRAPH.JUSTIFY, after=2)
        run(p, rot + ": ", bold=True)
        run(p, txt)

    secao("Experiência Profissional")
    for titulo, periodo, empresa, bs in v["experiencias"]:
        p = par(before=8)
        p.paragraph_format.tab_stops.add_tab_stop(RIGHT, WD_TAB_ALIGNMENT.RIGHT)
        run(p, titulo, bold=True)
        run(p, "\t" + periodo)
        run(par(), empresa, bold=True)
        for i, b in enumerate(bs):
            bp = doc.add_paragraph(style="List Bullet")
            bp.paragraph_format.space_after = Pt(1.5)
            bp.paragraph_format.space_before = Pt(3 if i == 0 else 0)
            bp.paragraph_format.left_indent = Inches(0.3)
            bp.paragraph_format.line_spacing = 1.05
            rich(bp, b)

    secao("Formação Acadêmica")
    p = par()
    p.paragraph_format.tab_stops.add_tab_stop(RIGHT, WD_TAB_ALIGNMENT.RIGHT)
    run(p, FORMACAO[0], bold=True)
    run(p, "\t" + FORMACAO[1])
    run(par(), FORMACAO[2])

    secao("Idiomas")
    for idioma, nivel in IDIOMAS:
        p = par()
        run(p, idioma, bold=True)
        run(p, " — " + nivel)

    doc.save(v["arquivo"] + ".docx")


if __name__ == "__main__":
    alvos = sys.argv[1:] or list(VERSOES)
    for nome in alvos:
        v = VERSOES[nome]
        gerar_html(v)
        gerar_docx(v)
        print(f"{nome}: {v['html']} + {v['arquivo']}.docx")
